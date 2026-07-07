from __future__ import annotations

import shutil
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
BUILD = ROOT / "examples/platformio/knob21/.pio/build/BOARD_VIEWE_UEDX48480021_MD80ET"
BOOT_APP0 = Path.home() / ".platformio/packages/framework-arduinoespressif32/tools/partitions/boot_app0.bin"
FLASH_TOOL_EXE = ROOT / "tools/flash_download_tool_3.9.5/flash_download_tool_3.9.3.exe"

OUT_DIR = ROOT / "field_delivery_package_20260707"
ZIP_PATH = ROOT / "digital_human_knob_field_delivery_package_20260707.zip"


TITLE = "数字人旋钮设备现场问题与烧录说明"


SECTIONS = [
    (
        "一、现场问题说明",
        [
            "现场安装设备时，由于开孔或固定角度存在偏差，设备外壳装上后屏幕方向看起来是歪的。这个问题可以通过软件补偿处理。",
            "如果偏差是 90°、180°、270° 这类整方向旋转，可以通过显示方向配置处理。",
            "如果偏差是几度的小角度倾斜，例如 3°、5°、8°，推荐在固件生成前统一反向旋转素材，再裁切成设备需要的 480×480 图片，然后重新生成 LVGL 图片数组并烧录。",
        ],
    ),
    (
        "二、推荐处理方案",
        [
            "推荐方案：预处理素材。把所有人物头像页和详情页统一按现场偏差角度进行反向旋转，补足边缘后裁切为 480×480，再重新生成固件。",
            "不推荐方案：设备运行时实时旋转全屏图片。虽然 LVGL 支持图片角度设置，但全屏图片实时旋转会增加 ESP32-S3 的显示负担，也可能出现边角露黑、刷新变慢等问题。",
            "现场如果后续确认了具体歪斜角度，只需要把角度告诉开发人员，即可重新生成素材和固件。",
        ],
    ),
    (
        "三、当前程序素材位置",
        [
            "原始 PNG 素材目录：examples/platformio/knob21/assets/digital_humans",
            "人物选择页头像：examples/platformio/knob21/assets/digital_humans/avatars",
            "按压选中后的详情页：examples/platformio/knob21/assets/digital_humans/responses",
            "真正编译进固件的 LVGL 图片数组：examples/platformio/knob21/lib/ui/src/images",
            "头像数组文件：ui_img_avatar_*.c",
            "详情页数组文件：ui_img_response_*.c",
        ],
    ),
    (
        "四、当前人物顺序",
        [
            "当前设备人物顺序为：诸葛亮、爱因斯坦、居里夫人、苏轼、孔子。",
            "第一个人物槽位程序符号名仍沿用旧名称 ui_img_avatar_newton / ui_img_response_newton，但实际内容已经替换为诸葛亮。",
            "设备按压选中后发出的 JSON 中，第一个人物 name 字段已经改为 zhuge_liang。",
        ],
    ),
    (
        "五、烧录包内容",
        [
            "firmware/bootloader.bin：引导程序",
            "firmware/partitions.bin：分区表",
            "firmware/boot_app0.bin：Arduino/ESP32 启动辅助文件",
            "firmware/firmware.bin：主程序固件",
            "scripts/flash_com3.bat：Windows 一键烧录脚本",
            "tools/flash_download_tool_3.9.3.exe：Espressif 图形化烧录工具",
            "docs/现场问题与烧录说明.docx：Word 说明文档",
            "docs/现场问题与烧录说明_UTF8.txt：UTF-8 BOM 文本说明，避免中文乱码",
        ],
    ),
    (
        "六、一键脚本烧录方式",
        [
            "1. 现场电脑安装 Python。",
            "2. 安装 esptool：python -m pip install esptool",
            "3. 使用 USB 线连接设备，确认 Windows 设备管理器中的串口号，例如 COM3。",
            "4. 双击 scripts/flash_com3.bat。",
            "5. 如果串口不是 COM3，例如是 COM5，则在命令行进入 scripts 目录后执行：flash_com3.bat COM5",
        ],
    ),
    (
        "七、Flash Download Tool 图形工具烧录方式",
        [
            "芯片选择：ESP32-S3",
            "模式选择：SPI Download",
            "SPI SPEED：80MHz 或 40MHz",
            "SPI MODE：DIO",
            "FLASH SIZE：16MB",
            "BAUD：921600；如果失败可改为 115200",
        ],
    ),
    (
        "八、图形工具文件地址填写",
        [
            "firmware/bootloader.bin    0x0000",
            "firmware/partitions.bin    0x8000",
            "firmware/boot_app0.bin     0xE000",
            "firmware/firmware.bin      0x10000",
        ],
    ),
    (
        "九、设备进入烧录模式",
        [
            "通常设备会自动进入烧录模式，不需要手动操作。",
            "如果连接失败，可以按住 BOOT 键或设备中间旋钮按压键，再点击 RESET/EN 或重新插拔 USB。",
            "看到工具开始 Connecting... 后松开 BOOT/旋钮按压键。",
        ],
    ),
    (
        "十、烧录成功判断",
        [
            "看到 Hash of data verified 表示写入校验通过。",
            "看到 SUCCESS 或 Leaving... Hard resetting 表示烧录流程完成，设备会自动重启。",
            "如果失败，优先检查串口号是否正确、USB 线是否支持数据传输、是否被其他串口工具占用，再尝试降低波特率到 115200。",
        ],
    ),
]


def set_run_font(run, name: str = "微软雅黑", size: int | None = None, bold: bool | None = None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def set_paragraph_font(paragraph, name: str = "微软雅黑", size: int = 10):
    for run in paragraph.runs:
        set_run_font(run, name=name, size=size)


def add_shaded_note(document: Document, text: str):
    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.left_indent = Cm(0.2)
    run = p.add_run(text)
    set_run_font(run, size=10, bold=True)
    run.font.color.rgb = RGBColor(120, 72, 0)
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "FFF3CD")
    p_pr.append(shd)


def build_docx(path: Path):
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(1.9)
    section.right_margin = Cm(1.9)

    styles = document.styles
    styles["Normal"].font.name = "微软雅黑"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    styles["Normal"].font.size = Pt(10.5)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run(TITLE)
    set_run_font(title_run, size=18, bold=True)
    title_run.font.color.rgb = RGBColor(24, 74, 124)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run("适用对象：现场安装、固件烧录、素材角度校正沟通")
    set_run_font(subtitle_run, size=10)
    subtitle_run.font.color.rgb = RGBColor(90, 90, 90)

    add_shaded_note(
        document,
        "交付说明：本压缩包已经包含固件、烧录脚本和图形化烧录工具。现场只需要解压后按本文档操作即可。",
    )

    for heading, paragraphs in SECTIONS:
        h = document.add_paragraph()
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(4)
        r = h.add_run(heading)
        set_run_font(r, size=13, bold=True)
        r.font.color.rgb = RGBColor(24, 74, 124)

        for text in paragraphs:
            p = document.add_paragraph(style=None)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            if text[:2].isdigit() or text.startswith("firmware/") or text.startswith("scripts/") or text.startswith("tools/") or text.startswith("docs/"):
                p.paragraph_format.left_indent = Cm(0.45)
            run = p.add_run(text)
            set_run_font(run, size=10.5)

    document.save(path)


def build_txt(path: Path):
    lines = [TITLE, "", "适用对象：现场安装、固件烧录、素材角度校正沟通", ""]
    for heading, paragraphs in SECTIONS:
        lines.append(heading)
        for paragraph in paragraphs:
            lines.append(f"  {paragraph}")
        lines.append("")
    path.write_text("\r\n".join(lines), encoding="utf-8-sig")


def copy_required_files():
    if OUT_DIR.exists():
        shutil.rmtree(OUT_DIR)
    (OUT_DIR / "docs").mkdir(parents=True)
    (OUT_DIR / "firmware").mkdir()
    (OUT_DIR / "scripts").mkdir()
    (OUT_DIR / "tools").mkdir()

    required = {
        BUILD / "bootloader.bin": OUT_DIR / "firmware/bootloader.bin",
        BUILD / "partitions.bin": OUT_DIR / "firmware/partitions.bin",
        BOOT_APP0: OUT_DIR / "firmware/boot_app0.bin",
        BUILD / "firmware.bin": OUT_DIR / "firmware/firmware.bin",
        FLASH_TOOL_EXE: OUT_DIR / "tools/flash_download_tool_3.9.3.exe",
    }
    for src, dst in required.items():
        if not src.exists():
            raise FileNotFoundError(src)
        shutil.copy2(src, dst)

    flash_bat = OUT_DIR / "scripts/flash_com3.bat"
    flash_bat.write_text(
        "@echo off\r\n"
        "setlocal\r\n"
        "set PORT=COM3\r\n"
        "if not \"%~1\"==\"\" set PORT=%~1\r\n"
        "cd /d \"%~dp0..\\firmware\"\r\n"
        "python -m esptool --chip esp32s3 --port %PORT% --baud 921600 --before default_reset --after hard_reset write_flash -z ^\r\n"
        "  0x0000 bootloader.bin ^\r\n"
        "  0x8000 partitions.bin ^\r\n"
        "  0xE000 boot_app0.bin ^\r\n"
        "  0x10000 firmware.bin\r\n"
        "pause\r\n",
        encoding="utf-8-sig",
    )


def zip_dir(src_dir: Path, zip_path: Path):
    if zip_path.exists():
        zip_path.unlink()
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
        for file in src_dir.rglob("*"):
            if file.is_file():
                zf.write(file, file.relative_to(src_dir.parent))


def main():
    copy_required_files()
    build_docx(OUT_DIR / "docs/现场问题与烧录说明.docx")
    build_txt(OUT_DIR / "docs/现场问题与烧录说明_UTF8.txt")
    zip_dir(OUT_DIR, ZIP_PATH)
    print(OUT_DIR)
    print(ZIP_PATH)


if __name__ == "__main__":
    main()
